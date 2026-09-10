import importlib
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.defined_name import DefinedName

from conftest import make_landscape_template, make_minimal_template, make_mismatched_columns_template

excel_coords = importlib.import_module("app.import.excel_coords")
cell_top_left_pt = excel_coords.cell_top_left_pt
col_width_to_pt = excel_coords.col_width_to_pt
merged_range_width_pt = excel_coords.merged_range_width_pt


def _parse_template(xlsx_bytes: bytes):
    module = importlib.import_module("app.import.template_importer")
    return module.parse_template(xlsx_bytes, name="chung-tu", version="1")


def _change_name(xlsx_bytes: bytes, name: str, target: str) -> bytes:
    workbook = load_workbook(BytesIO(xlsx_bytes))
    del workbook.defined_names[name]
    workbook.defined_names.add(DefinedName(name, attr_text=target))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _save_workbook(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _make_two_row_header_template(*, vertical_merge_first_column: bool = False) -> bytes:
    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]

    sheet.unmerge_cells("B5:C5")
    sheet.unmerge_cells("E5:F5")
    sheet.unmerge_cells("B6:C6")
    sheet.unmerge_cells("E6:F6")

    if vertical_merge_first_column:
        sheet.merge_cells("A5:A6")
        sheet["A5"] = "STT"
    else:
        sheet["A5"] = "Nhóm"
        sheet["A6"] = "STT"

    sheet["B5"] = "Hàng hóa"
    sheet["D5"] = "Định lượng"
    sheet["E5"] = "Ghi nhận"
    sheet["B6"] = "Tên hàng"
    sheet["D6"] = "ĐVT"
    sheet["E6"] = "Số lượng"
    sheet["G6"] = "Ghi chú"
    sheet.merge_cells("B6:C6")
    sheet.merge_cells("E6:F6")
    sheet.merge_cells("B7:C7")
    sheet.merge_cells("E7:F7")
    sheet.row_dimensions[5].height = 20
    sheet.row_dimensions[6].height = 24

    del workbook.defined_names["TABLE_HEADER"]
    del workbook.defined_names["TABLE_DATA_ROW"]
    workbook.defined_names.add(
        DefinedName("TABLE_HEADER", attr_text="'ChungTu'!$A$5:$G$6")
    )
    workbook.defined_names.add(
        DefinedName("TABLE_DATA_ROW", attr_text="'ChungTu'!$A$7:$G$7")
    )
    return _save_workbook(workbook)


def _make_pnk_like_two_row_header() -> bytes:
    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]

    for merged_range in ("B5:C5", "E5:F5", "B6:C6", "E6:F6"):
        sheet.unmerge_cells(merged_range)

    sheet["A5"] = "TT"
    sheet["B5"] = "Tên, nhãn hiệu, quy cách, phẩm chất vật tư, dụng cụ, sản phẩm, hàng hóa"
    sheet["C5"] = "ĐVT"
    sheet["D5"] = "Số lượng"
    sheet["D6"] = "Yêu cầu"
    sheet["E6"] = "Thực nhập"
    sheet["F5"] = "Đơn giá"
    sheet["G5"] = "Thành tiền"
    sheet.merge_cells("A5:A6")
    sheet.merge_cells("B5:B6")
    sheet.merge_cells("C5:C6")
    sheet.merge_cells("D5:E5")
    sheet.merge_cells("F5:F6")
    sheet.merge_cells("G5:G6")
    sheet.row_dimensions[5].height = 20
    sheet.row_dimensions[6].height = 24

    del workbook.defined_names["TABLE_HEADER"]
    del workbook.defined_names["TABLE_DATA_ROW"]
    workbook.defined_names.add(
        DefinedName("TABLE_HEADER", attr_text="'ChungTu'!$A$5:$G$6")
    )
    workbook.defined_names.add(
        DefinedName("TABLE_DATA_ROW", attr_text="'ChungTu'!$A$7:$G$7")
    )
    return _save_workbook(workbook)


def test_parse_minimal_template():
    metadata = _parse_template(make_minimal_template())

    assert metadata.name == "chung-tu"
    assert metadata.version == "1"
    assert metadata.table.sheet_name == "ChungTu"
    assert metadata.table.header_row_range == "A5:G5"
    assert metadata.table.data_row_template == "A6:G6"
    assert [field.field_name for field in metadata.fields] == [
        "don_vi",
        "ngay_thang",
    ]
    assert [column.key for column in metadata.table.columns] == [
        "stt",
        "ten_hang",
        "dvt",
        "so_luong",
        "ghi_chu",
    ]
    assert metadata.static_cells
    assert any(cell.layer == "header" for cell in metadata.static_cells)
    assert metadata.page.static_block_height_pt > 0
    from app.template.page_size import page_dimensions

    page_width, _ = page_dimensions(metadata.page)
    usable = page_width - metadata.page.margin_left - metadata.page.margin_right
    assert abs(sum(column.width_pt for column in metadata.table.columns) - usable) < 0.5
    assert metadata.table.row_style["font"]["size"] == 11.0
    assert metadata.table.header_height_pt == 24
    assert metadata.table.signature_block_height_pt == 80


def test_table_header_two_rows_imports_keys_from_bottom_row():
    metadata = _parse_template(_make_two_row_header_template())

    assert metadata.table.header_row_range == "A5:G6"
    assert metadata.table.data_row_template == "A7:G7"
    assert [column.key for column in metadata.table.columns] == [
        "stt",
        "ten_hang",
        "dvt",
        "so_luong",
        "ghi_chu",
    ]
    assert metadata.table.header_height_pt == 44
    header_rows = {cell.row for cell in metadata.static_cells if cell.layer == "header"}
    assert header_rows == {5, 6}
    assert any(
        cell.layer == "header" and cell.value == "Hàng hóa"
        for cell in metadata.static_cells
    )


def test_table_header_vertical_merge_reads_title_from_merge_origin():
    metadata = _parse_template(_make_two_row_header_template(vertical_merge_first_column=True))

    assert metadata.table.columns[0].title == "STT"
    assert metadata.table.columns[0].key == "stt"


def test_header_vertical_merge_static_cell_uses_full_height():
    metadata = _parse_template(_make_pnk_like_two_row_header())

    tt = next(
        cell
        for cell in metadata.static_cells
        if cell.layer == "header" and cell.value == "TT"
    )
    so_luong = next(
        cell
        for cell in metadata.static_cells
        if cell.layer == "header" and cell.value == "Số lượng"
    )

    assert tt.height_pt == pytest.approx(44.0)
    assert so_luong.width_pt > tt.width_pt


def test_table_header_rejects_upper_row_merge_outside_bounds():
    errors = importlib.import_module("app.import.errors")
    workbook = load_workbook(BytesIO(_make_two_row_header_template()))
    sheet = workbook["ChungTu"]
    sheet.merge_cells("G5:H5")

    with pytest.raises(errors.TemplateValidationError, match="vượt ngoài biên"):
        _parse_template(_save_workbook(workbook))


def test_parse_sheet_scoped_named_ranges():
    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    for name, defined_name in list(workbook.defined_names.items()):
        cell_range = defined_name.attr_text.split("!", 1)[1]
        del workbook.defined_names[name]
        sheet.defined_names.add(DefinedName(name, attr_text=cell_range))
    output = BytesIO()
    workbook.save(output)

    metadata = _parse_template(output.getvalue())

    assert metadata.table.header_row_range == "A5:G5"
    assert {field.field_name for field in metadata.fields} == {"don_vi", "ngay_thang"}


def test_invalid_signature_height_uses_default():
    workbook = load_workbook(BytesIO(make_minimal_template()))
    workbook.defined_names.add(
        DefinedName("TABLE_SIGNATURE", attr_text="'ChungTu'!$A$100")
    )
    output = BytesIO()
    workbook.save(output)

    metadata = _parse_template(output.getvalue())

    assert metadata.table.signature_block_height_pt == 80


def test_rejects_mismatched_column_count():
    errors = importlib.import_module("app.import.errors")

    with pytest.raises(errors.TemplateValidationError, match=r"5.*4"):
        _parse_template(make_mismatched_columns_template())


def test_rejects_missing_required_named_range():
    errors = importlib.import_module("app.import.errors")
    workbook = load_workbook(BytesIO(make_minimal_template()))
    del workbook.defined_names["TABLE_HEADER"]
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(errors.TemplateValidationError, match="TABLE_HEADER"):
        _parse_template(output.getvalue())


def test_rejects_table_ranges_on_different_sheets():
    errors = importlib.import_module("app.import.errors")
    workbook = load_workbook(BytesIO(make_minimal_template()))
    workbook.create_sheet("Other")
    output = BytesIO()
    workbook.save(output)
    changed = _change_name(
        output.getvalue(), "TABLE_DATA_ROW", "'Other'!$A$6:$G$6"
    )

    with pytest.raises(errors.TemplateValidationError, match="cùng sheet"):
        _parse_template(changed)


def test_rejects_misaligned_column_bounds():
    errors = importlib.import_module("app.import.errors")
    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    sheet.unmerge_cells("B6:C6")
    sheet.unmerge_cells("E6:F6")
    sheet.merge_cells("C6:D6")
    sheet.merge_cells("F6:G6")
    del workbook.defined_names["TABLE_DATA_ROW"]
    workbook.defined_names.add(
        DefinedName("TABLE_DATA_ROW", attr_text="'ChungTu'!$B$6:$H$6")
    )
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(errors.TemplateValidationError, match="biên cột"):
        _parse_template(output.getvalue())


def test_rejects_invalid_field_name():
    errors = importlib.import_module("app.import.errors")
    workbook = load_workbook(BytesIO(make_minimal_template()))
    workbook.defined_names.add(
        DefinedName("FIELD_Don-Vi", attr_text="'ChungTu'!$A$2")
    )
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(errors.TemplateValidationError, match="FIELD_Don-Vi"):
        _parse_template(output.getvalue())


def test_parse_accepts_nl_field_named_ranges():
    workbook = load_workbook(BytesIO(make_minimal_template()))
    workbook.defined_names.add(
        DefinedName("NL_FIELD_can_cu_pnk", attr_text="'ChungTu'!$F$2")
    )
    output = BytesIO()
    workbook.save(output)

    metadata = _parse_template(output.getvalue())

    nl_field = next(field for field in metadata.fields if field.field_name == "can_cu_pnk")
    assert nl_field.sheet_name == "ChungTu"
    assert nl_field.cell_ref == "F2"
    assert nl_field.named_range == "NL_FIELD_can_cu_pnk"


def test_page_meta_reads_margin_and_orientation_from_file():
    """_page_meta phải đọc giá trị thật từ ws.page_margins/page_setup, không rơi về default."""
    MARGIN_TOP_IN = 1.0   # inch
    MARGIN_LEFT_IN = 0.75  # inch
    metadata = _parse_template(
        make_landscape_template(margin_top_in=MARGIN_TOP_IN, margin_left_in=MARGIN_LEFT_IN)
    )
    page = metadata.page
    assert page.orientation == "landscape"
    assert page.page_size == "A4"
    assert abs(page.margin_top - MARGIN_TOP_IN * 72) < 1.0   # inch → pt, tolerance 1pt
    assert abs(page.margin_left - MARGIN_LEFT_IN * 72) < 1.0


def test_rejects_full_row_named_range():
    errors = importlib.import_module("app.import.errors")
    changed = _change_name(make_minimal_template(), "TABLE_HEADER", "'ChungTu'!$5:$5")

    with pytest.raises(errors.TemplateValidationError, match="TABLE_HEADER"):
        _parse_template(changed)


def test_table_data_row_still_rejects_multi_row():
    errors = importlib.import_module("app.import.errors")
    changed = _change_name(
        make_minimal_template(), "TABLE_DATA_ROW", "'ChungTu'!$A$6:$G$7"
    )

    with pytest.raises(errors.TemplateValidationError, match="một dòng"):
        _parse_template(changed)


def test_column_align_h_comes_from_table_data_row_not_header():
    """align_h từng cột lấy từ TABLE_DATA_ROW (không lấy từ tiêu đề)."""
    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    # Header center everywhere — nếu đọc nhầm header sẽ ra center.
    for col in range(1, 8):
        sheet.cell(5, col).alignment = Alignment(horizontal="center")
    sheet["A6"].alignment = Alignment(horizontal="center")
    sheet["B6"].alignment = Alignment(horizontal="left")
    sheet["D6"].alignment = Alignment(horizontal="right")
    sheet["E6"].alignment = Alignment(horizontal="right")
    sheet["G6"].alignment = Alignment(horizontal="left")

    metadata = _parse_template(_save_workbook(workbook))
    by_key = {column.key: column.align_h for column in metadata.table.columns}
    assert by_key["stt"] == "center"
    assert by_key["ten_hang"] == "left"
    assert by_key["dvt"] == "right"
    assert by_key["so_luong"] == "right"
    assert by_key["ghi_chu"] == "left"


def test_field_geometry_uses_enclosing_merge():
    """FIELD_* on one cell inside a larger merge uses full merge box."""
    from app.template.page_size import page_dimensions

    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    sheet.merge_cells("B2:D2")
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center")
    del workbook.defined_names["FIELD_ngay_thang"]
    workbook.defined_names.add(
        DefinedName("FIELD_ngay_thang_nam", attr_text="'ChungTu'!$B$2")
    )
    output = BytesIO()
    workbook.save(output)
    xlsx_bytes = output.getvalue()

    metadata = _parse_template(xlsx_bytes)
    field = next(f for f in metadata.fields if f.field_name == "ngay_thang_nam")

    parsed_wb = load_workbook(BytesIO(xlsx_bytes))
    parsed_sheet = parsed_wb["ChungTu"]
    raw_merge_width = merged_range_width_pt(parsed_sheet, 2, 4)
    raw_single_width = col_width_to_pt(parsed_sheet.column_dimensions["B"].width)
    header_groups = [(1, 1), (2, 3), (4, 4), (5, 6), (7, 7)]
    raw_table_total = sum(
        merged_range_width_pt(parsed_sheet, min_col, max_col)
        for min_col, max_col in header_groups
    )
    page_width, _ = page_dimensions(metadata.page)
    usable = page_width - metadata.page.margin_left - metadata.page.margin_right
    scale = usable / raw_table_total

    assert field.width_pt == pytest.approx(raw_merge_width * scale)
    assert field.width_pt != pytest.approx(raw_single_width * scale)

    expected_x, _ = cell_top_left_pt(
        parsed_sheet,
        2,
        2,
        page_height=page_dimensions(metadata.page)[1],
        margin_top=metadata.page.margin_top,
        margin_left=metadata.page.margin_left,
    )
    table_left_x, _ = cell_top_left_pt(
        parsed_sheet,
        5,
        1,
        page_height=page_dimensions(metadata.page)[1],
        margin_top=metadata.page.margin_top,
        margin_left=metadata.page.margin_left,
    )
    if abs(scale - 1.0) >= 0.001:
        expected_x = table_left_x + (expected_x - table_left_x) * scale
    assert field.x == pytest.approx(expected_x)
    assert field.align["h"] == "center"


def test_interior_field_merge_skips_static_cell():
    """FIELD_* on interior merge cell skips full merge in static_cells."""
    from app.template.page_size import page_dimensions

    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    sheet.merge_cells("B2:D2")
    sheet["B2"] = "Ngày tháng năm"
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center")
    del workbook.defined_names["FIELD_ngay_thang"]
    workbook.defined_names.add(
        DefinedName("FIELD_ngay_thang_nam", attr_text="'ChungTu'!$C$2")
    )
    output = BytesIO()
    workbook.save(output)
    xlsx_bytes = output.getvalue()

    metadata = _parse_template(xlsx_bytes)
    field = next(f for f in metadata.fields if f.field_name == "ngay_thang_nam")

    parsed_wb = load_workbook(BytesIO(xlsx_bytes))
    parsed_sheet = parsed_wb["ChungTu"]
    raw_merge_width = merged_range_width_pt(parsed_sheet, 2, 4)
    header_groups = [(1, 1), (2, 3), (4, 4), (5, 6), (7, 7)]
    raw_table_total = sum(
        merged_range_width_pt(parsed_sheet, min_col, max_col)
        for min_col, max_col in header_groups
    )
    page_width, _ = page_dimensions(metadata.page)
    usable = page_width - metadata.page.margin_left - metadata.page.margin_right
    scale = usable / raw_table_total

    assert field.width_pt == pytest.approx(raw_merge_width * scale)

    field_right = field.x + field.width_pt
    overlapping = [
        cell
        for cell in metadata.static_cells
        if cell.row == 2
        and cell.x < field_right
        and cell.x + cell.width_pt > field.x
    ]
    assert overlapping == []


def test_interior_nl_field_merge_skips_static_cell():
    """NL_FIELD_* on interior merge cell skips full merge in static_cells."""
    from app.template.page_size import page_dimensions

    workbook = load_workbook(BytesIO(make_minimal_template()))
    sheet = workbook["ChungTu"]
    sheet.merge_cells("E2:G2")
    sheet["E2"] = "Căn cứ"
    sheet["E2"].alignment = Alignment(horizontal="center", vertical="center")
    workbook.defined_names.add(
        DefinedName("NL_FIELD_can_cu_pnk", attr_text="'ChungTu'!$F$2")
    )
    output = BytesIO()
    workbook.save(output)
    xlsx_bytes = output.getvalue()

    metadata = _parse_template(xlsx_bytes)
    field = next(f for f in metadata.fields if f.field_name == "can_cu_pnk")

    parsed_wb = load_workbook(BytesIO(xlsx_bytes))
    parsed_sheet = parsed_wb["ChungTu"]
    raw_merge_width = merged_range_width_pt(parsed_sheet, 5, 7)
    header_groups = [(1, 1), (2, 3), (4, 4), (5, 6), (7, 7)]
    raw_table_total = sum(
        merged_range_width_pt(parsed_sheet, min_col, max_col)
        for min_col, max_col in header_groups
    )
    page_width, _ = page_dimensions(metadata.page)
    usable = page_width - metadata.page.margin_left - metadata.page.margin_right
    scale = usable / raw_table_total

    assert field.width_pt == pytest.approx(raw_merge_width * scale)
    assert field.named_range == "NL_FIELD_can_cu_pnk"

    field_right = field.x + field.width_pt
    overlapping = [
        cell
        for cell in metadata.static_cells
        if cell.row == 2
        and cell.x < field_right
        and cell.x + cell.width_pt > field.x
    ]
    assert overlapping == []


def test_named_range_tokenizer_error_maps_to_validation_error():
    """Excel sheet-copy names like `01 (2)` can produce attr_text openpyxl cannot tokenize."""
    errors = importlib.import_module("app.import.errors")

    workbook = load_workbook(BytesIO(make_minimal_template()))
    workbook.create_sheet("01 (2)")
    workbook.defined_names.add(
        DefinedName("FIELD_so_phieu", attr_text="=$'01 (2)'.$A$5")
    )
    with pytest.raises(errors.TemplateValidationError, match="không đọc được"):
        _parse_template(_save_workbook(workbook))
