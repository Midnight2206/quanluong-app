"""So bản in ngày 01 và 03 với PDF mẫu: vị trí cột và số dòng mỗi tờ."""

from __future__ import annotations

import os
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

os.environ.setdefault("DOCUMENT_SERVICE_KEY", "test-key")

from app.render.kitchen_ledger_book import header_merge_boxes, render_kitchen_ledger_book

ROOT = Path("/workspace/docs/print-templates")
MAU = ROOT / "mau"
REF = ROOT / "ref"
ART = Path("/opt/cursor/artifacts/print-compare")

LEFT_X = [31.2, 55.7, 215.0, 241.3, 283.8, 315.4, 374.0, 409.3, 444.6, 477.4, 510.3, 563.5]
RIGHT_X = [31.2, 64.0, 116.0, 149.5, 202.7, 234.3, 267.2, 298.8, 350.9, 560.2]
KEYS = [
    "stt", "name", "uom", "price", "openingQty", "openingAmount",
    "inG", "inH", "inI", "inJ", "inAmount", "totalInQty", "totalInAmount",
    "mealQty", "mealAmount", "otherQty", "otherAmount", "totalOutQty", "totalOutAmount",
]


def _qty(value):
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value.strip()
    if abs(float(value)) < 1e-9:
        return ""
    text = f"{float(value):.4f}".rstrip("0").rstrip(".")
    if "." in text:
        whole, frac = text.split(".", 1)
        return f"{whole},{frac}"
    return text


def _money(value):
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value.strip()
    if abs(float(value)) < 0.5:
        return ""
    return f"{int(round(float(value))):,}".replace(",", ".")


def _cell(ws, row, col, money=False):
    return (_money if money else _qty)(ws.cell(row, col).value)


def _row(ws, row):
    money_cols = {4, 6, 11, 13, 15, 17, 19}
    out = {}
    for index, key in enumerate(KEYS, start=1):
        out[key] = _cell(ws, row, index, money=index in money_cols)
    return out


def _finance(ws):
    lines = []
    for row in range(7, 35):
        label = ws.cell(row, 20).value
        value = ws.cell(row, 21).value
        if label is None and value is None:
            continue
        label = "" if label is None else str(label).strip()
        if isinstance(value, (int, float)):
            shown = _money(value)
        else:
            shown = "" if value is None else str(value).strip()
        lines.append({"label": label, "value": shown})
    return lines


def _day(path, date_line):
    ws = load_workbook(path, data_only=True).active
    sheets = []
    current = []
    carry_out = None
    for row in range(7, ws.max_row + 1):
        name = ws.cell(row, 2).value
        if name is None:
            continue
        label = str(name).strip()
        if label == "Cộng sang trang":
            carry_out = _row(ws, row)
            sheets.append({"rows": current, "carryOut": carry_out, "carryIn": None, "showTotal": False, "showSignature": False})
            current = []
            continue
        if label == "Mang sang":
            continue
        if label == "Cộng":
            sheets.append({
                "rows": current,
                "carryIn": carry_out,
                "carryOut": None,
                "showTotal": True,
                "showSignature": True,
                "totals": _row(ws, row),
            })
            current = []
            continue
        if label in ("QUẢN LÝ ĐƠN VỊ",):
            break
        if ws.cell(row, 1).value is None and label != "Cộng":
            continue
        current.append(_row(ws, row))
    if current:
        sheets.append({
            "rows": current,
            "carryIn": carry_out,
            "carryOut": None,
            "showTotal": True,
            "showSignature": True,
            "totals": {},
        })
    return {
        "dateLine": date_line,
        "draft": False,
        "financeLines": _finance(ws),
        "sheets": sheets,
    }


def _book():
    return {
        "cover": {
            "commanderRankName": "Trung tá Nguyễn Minh Tuấn",
            "keeper": "1//CN Nguyễn Văn Khánh",
            "dateLine": "Ngày 01 tháng 10 năm 2026",
            "draftDays": [],
        },
        "signatures": {
            "quan_ly": "1// CN Nguyễn Văn Khánh",
            "truc_ban": "",
            "chi_huy": "2//  Nguyễn Minh Tuấn",
        },
        "days": [
            _day(MAU / "cktc_01.xlsx", "Ngày 01 tháng 10 năm 2026"),
            _day(MAU / "cktc_03.xlsx", "Ngày 03 tháng 10 năm 2026"),
        ],
        "tail": None,
    }


def _verticals(page):
    xs = []
    for line in page.get_drawings():
        for item in line.get("items", []):
            if item[0] != "l":
                continue
            p1, p2 = item[1], item[2]
            if abs(p1.x - p2.x) < 0.4 and abs(p1.y - p2.y) > 8:
                xs.append(round(p1.x, 1))
    return sorted(set(xs))


def _horizontals(page):
    ys = []
    for line in page.get_drawings():
        for item in line.get("items", []):
            if item[0] != "l":
                continue
            p1, p2 = item[1], item[2]
            if abs(p1.y - p2.y) < 0.4 and abs(p1.x - p2.x) > 8:
                ys.append(round(p1.y, 1))
    return sorted(set(ys))


def _segments(page):
    found = []
    for drawing in page.get_drawings():
        for item in drawing.get("items", []):
            if item[0] != "l":
                continue
            p1, p2 = item[1], item[2]
            found.append((p1.x, p1.y, p2.x, p2.y))
    return found


def _line_cuts_box(seg, box, eps=0.6):
    x0, y0, x1, y1 = box
    ax, ay, bx, by = seg
    if abs(ax - bx) < 0.4:
        x = (ax + bx) / 2
        if not (x0 + eps < x < x1 - eps):
            return False
        top, bot = min(ay, by), max(ay, by)
        return min(bot, y1 - eps) - max(top, y0 + eps) > 0.4
    if abs(ay - by) < 0.4:
        y = (ay + by) / 2
        if not (y0 + eps < y < y1 - eps):
            return False
        left, right = min(ax, bx), max(ax, bx)
        return min(right, x1 - eps) - max(left, x0 + eps) > 0.4
    return False


def test_merged_header_cells_have_no_inner_line():
    pdf = render_kitchen_ledger_book(_book())
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    # Tờ đầu ngày 01: trang trái rồi trang phải.
    pages = {"left": doc[1], "right": doc[2]}
    named = {
        "left": [
            "Ngày trước chuyển qua",
            "Nhập trong ngày",
            "Số lượng nhập",
            "Trên BĐ",
            "Đơn vị BĐ",
            "Mua TT trên",
            "TG CB trên",
            "Mua TT đơn vị",
            "TG CB đơn vị",
            "TT",
            "Tên LTTP-Chất đốt",
        ],
        "right": [
            "Cộng nhập",
            "Xuất ăn trong ngày",
            "Xuất khác",
            "Cộng xuất",
            "CÔNG KHAI TÀI CHÍNH",
            "Số lượng xuất ăn",
            "Thành tiền xuất ăn",
        ],
    }
    for side, page in pages.items():
        segs = _segments(page)
        boxes = {name: box for name, *box in header_merge_boxes(side)}
        for name in named[side]:
            box = boxes[name]
            cut = [seg for seg in segs if _line_cuts_box(seg, box)]
            assert not cut, (side, name, box, cut[:3])
    doc.close()


def test_sheet_breaks_match_sample_days():
    book = _book()
    day01, day03 = book["days"]
    assert [len(sheet["rows"]) for sheet in day01["sheets"]] == [43, 2]
    assert [len(sheet["rows"]) for sheet in day03["sheets"]] == [36, 2]


def test_column_lines_match_sample_within_half_point():
    pdf = render_kitchen_ledger_book(_book())
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    # Trang 0 là bìa. Ngày 01: trái tờ 1, phải tờ 1.
    left = _verticals(doc[1])
    right = _verticals(doc[2])
    for expected in LEFT_X:
        assert min(abs(expected - actual) for actual in left) <= 0.5, (expected, left)
    for expected in RIGHT_X:
        assert min(abs(expected - actual) for actual in right) <= 0.5, (expected, right)
    left_y = _horizontals(doc[1])
    for expected in (57.7, 88.9, 103.5, 118.0, 147.1):
        assert min(abs(expected - actual) for actual in left_y) <= 0.6, (expected, left_y[:8])
    assert all(abs(x - 445.9) > 2 for x in right)
    next_top = _horizontals(doc[3])
    assert min(abs(43.1 - actual) for actual in next_top) <= 0.6, next_top[:4]
    doc.close()


@pytest.mark.skipif(not (REF / "snx-002.png").is_file(), reason="thiếu ảnh mẫu")
def test_raster_difference_at_110_dpi():
    pdf = render_kitchen_ledger_book(_book())
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    ART.mkdir(parents=True, exist_ok=True)
    pairs = [
        (0, REF / "snx-001.png", "cover", 0.08),
        (1, REF / "snx-002.png", "d01-left-1", 0.16),
        (2, REF / "snx-003.png", "d01-right-1", 0.16),
        (3, REF / "snx-004.png", "d01-left-2", 0.08),
        (4, REF / "snx-005.png", "d01-right-2", 0.08),
        (5, REF / "snx-day03-010.png", "d03-left-1", 0.16),
        (6, REF / "snx-day03-011.png", "d03-right-1", 0.16),
        (7, REF / "snx-day03-012.png", "d03-left-2", 0.08),
        (8, REF / "snx-day03-013.png", "d03-right-2", 0.08),
    ]
    scores = []
    for index, ref_path, name, limit in pairs:
        if not ref_path.is_file():
            continue
        pix = doc[index].get_pixmap(dpi=110, colorspace=pymupdf.csGRAY, alpha=False)
        pix.save(ART / f"after-{name}.png")
        ref = pymupdf.Pixmap(str(ref_path))
        if ref.alpha:
            ref = pymupdf.Pixmap(ref, 0)
        if ref.n != 1:
            ref = pymupdf.Pixmap(pymupdf.csGRAY, ref)
        ref.save(ART / f"before-{name}.png")
        width = min(pix.w, ref.w)
        height = min(pix.h, ref.h)
        differ = 0
        total = width * height
        got = pix.samples
        src = ref.samples
        gw, gh = pix.w, pix.h
        rw = ref.w
        for y in range(height):
            for x in range(width):
                a = got[y * gw + x]
                b = src[y * rw + x]
                if abs(a - b) > 48:
                    differ += 1
        ratio = differ / total
        scores.append((name, round(ratio, 4), limit))
    doc.close()
    (ART / "scores.txt").write_text("\n".join(f"{name} {ratio}" for name, ratio, _limit in scores), encoding="utf-8")
    # Font Tinos khác Times New Roman nên nét chữ lệch. Khung lệch thì tỉ lệ nhảy cao hơn ngưỡng.
    for name, ratio, limit in scores:
        assert ratio < limit, (name, ratio, scores)
