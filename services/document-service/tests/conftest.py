from io import BytesIO

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


def _add_name(workbook, name: str, sheet_name: str, cell_range: str) -> None:
    workbook.defined_names.add(
        DefinedName(name, attr_text=f"'{sheet_name}'!{cell_range}")
    )


def _template_bytes(*, data_range: str = "$A$6:$G$6") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ChungTu"

    sheet["A1"] = "Đơn vị"
    sheet["G2"] = "Ngày tháng"
    _add_name(workbook, "FIELD_don_vi", sheet.title, "$A$1")
    _add_name(workbook, "FIELD_ngay_thang", sheet.title, "$G$2")

    headers = ("STT", "Tên hàng", "ĐVT", "Số lượng", "Ghi chú")
    starts = ("A5", "B5", "D5", "E5", "G5")
    for cell_ref, title in zip(starts, headers):
        sheet[cell_ref] = title
    sheet.merge_cells("B5:C5")
    sheet.merge_cells("E5:F5")
    sheet.merge_cells("B6:C6")
    sheet.merge_cells("E6:F6")

    for column, width in zip("ABCDEFG", (5, 8, 9, 7, 6, 4, 12)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[5].height = 24

    _add_name(workbook, "TABLE_HEADER", sheet.title, "$A$5:$G$5")
    _add_name(workbook, "TABLE_DATA_ROW", sheet.title, data_range)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_minimal_template() -> bytes:
    return _template_bytes()


def make_mismatched_columns_template() -> bytes:
    return _template_bytes(data_range="$A$6:$F$6")
