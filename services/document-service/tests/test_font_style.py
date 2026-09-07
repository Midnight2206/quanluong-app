from openpyxl import Workbook
from openpyxl.styles import Font

from app.render.font_style import font_dict_from_cell, font_size_from_style, render_font


def test_font_dict_from_cell_preserves_size():
    wb = Workbook()
    ws = wb.active
    ws["A1"].font = Font(size=13, bold=True)
    payload = font_dict_from_cell(ws["A1"])
    assert payload["size"] == 13.0
    assert payload["bold"] is True
    name, size = render_font(payload)
    assert size == 13.0


def test_font_size_from_style_defaults_to_excel_11():
    assert font_size_from_style(None) == 11.0
    assert font_size_from_style({"size": None}) == 11.0
    assert font_size_from_style({"size": 9.5}) == 9.5
