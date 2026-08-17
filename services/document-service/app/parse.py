from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook


def _normalize_cell(value):
    if value is None or value == "":
        return None
    return value


def parse_xlsx(data: bytes, sheet: str | None = None) -> dict:
    if not data:
        raise ValueError("Dữ liệu Excel trống hoặc không hợp lệ")

    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValueError("Dữ liệu Excel trống hoặc không hợp lệ") from exc
    sheets = wb.sheetnames
    if not sheets:
        raise ValueError("Workbook Excel không có sheet nào")

    if sheet is not None:
        if sheet not in sheets:
            raise ValueError(f"Không tìm thấy sheet '{sheet}'")
        active = sheet
    else:
        active = sheets[0]

    ws = wb[active]
    max_col = ws.max_column or 0

    rows = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [_normalize_cell(v) for v in row]
        if max_col > len(cells):
            cells.extend([None] * (max_col - len(cells)))
        rows.append({"r": r_idx, "c": cells})

    while rows and all(c is None for c in rows[-1]["c"]):
        rows.pop()

    return {"sheet": active, "sheets": list(sheets), "rows": rows}
